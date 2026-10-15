# -*- coding: utf-8 -*-
"""生成周报 Word 文档。"""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt

DST = r"E:\BaiduNetdiskDownload\LINKOA\本周工作汇总_20261014.docx"

doc = Document()

# 默认字体
style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_h(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "微软雅黑"
        run.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_p(text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "微软雅黑"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    return p


# 标题
title = doc.add_heading("本周工作汇总", level=0)
for run in title.runs:
    run.font.name = "微软雅黑"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

add_p("周期：2026.10.08 – 2026.10.14", bold=True)

# 一、安全加固
add_h("一、安全加固", level=1)
items = [
    "密码存储升级：从明文改为 bcrypt 哈希加密存储，历史密码首次登录自动平滑升级；数据库即使泄露也无法反推原始密码。",
    "接口权限收口：用户列表接口不再返回密码字段；附件上传增加扩展名黑名单（拦截 html/svg/js 等可执行文件上传）；附件下载与列表增加客户单位数据隔离，避免客户跨公司查看附件。",
    "登录态安全：JWT 增加版本号机制，修改密码后旧 token 立即失效；JWT 密钥仍为默认值时启动打印安全警告，防止部署漏配。",
    "路径穿越修复：前端路由回退接口增加 realpath 校验，避免被构造特殊请求读到项目目录外文件。",
]
for it in items:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(it)
    run.font.name = "微软雅黑"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

# 二、性能与访问优化
add_h("二、性能与访问优化", level=1)
items = [
    "全站 GZip 压缩：HTML/JS/CSS 等文本资源自动压缩，首屏加载流量减少 60%-80%。",
    "静态资源长缓存：前端打包文件带内容 hash，配置一年强缓存，二次访问不再重复下载。",
    "数据库查询优化：列表分页 count 改 func.count()，避免全量 id 拉进内存；Excel 导出改流式查询，项目数据量大时不再 OOM。",
    "AI 助手加速：知识库查询加 5 分钟内存缓存，重复提问不再每次全量重算项目数据。",
]
for it in items:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(it)
    run.font.name = "微软雅黑"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

# 三、业务功能迭代
add_h("三、业务功能迭代", level=1)
items = [
    "项目卡片改造：从静态三时间（创建/开始/截止）改为动态多环节模式，管理员可按需追加环节（如“IP设计截止”→“打样制作”），客户可看到当前环节倒计时及每个环节实际延误天数。",
    "客户审核节点：新增独立节点类型，管理员可配置审核内容与选项；客户提交后不可修改，管理员有权增删改查；客户可查看自己提交的内容和上传的图片。",
    "BUG 修复：修复节点时间时区显示错误、项目卡片文案逻辑、客户审核页服务器内部错误、图片上传失败等问题。",
]
for it in items:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(it)
    run.font.name = "微软雅黑"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

# 四、下周计划
add_h("四、下周计划", level=1)
items = [
    "生产环境部署前完成数据迁移与回归测试。",
    "继续优化客户门户体验：审核待办高亮、消息推送。",
]
for it in items:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(it)
    run.font.name = "微软雅黑"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

doc.save(DST)
print("OK:", DST)
