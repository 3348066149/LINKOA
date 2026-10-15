# -*- coding: utf-8 -*-
"""把 REVIEW.md 转成结构化 Word 文档（REVIEW.docx）。"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(r"E:\BaiduNetdiskDownload\LINKOA")
SRC = ROOT / "REVIEW.md"
DST = ROOT / "REVIEW.docx"

doc = Document()

# 全局默认字体
style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def set_cell_font(cell, bold=False, size=9):
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.name = "微软雅黑"
            run.font.size = Pt(size)
            run.bold = bold
            run.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_rich_paragraph(text, style_name=None, bullet=False):
    """处理 **bold** 行内标记。"""
    p = doc.add_paragraph(style=style_name) if style_name else doc.add_paragraph()
    if bullet:
        p.style = doc.styles["List Bullet"]
    # 切分 **...**
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            run = p.add_run(part)
    return p


def flush_table(rows):
    """rows: list[list[str]]，第一行是表头。"""
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = val.strip()
            set_cell_font(cell, bold=(i == 0), size=9)
    doc.add_paragraph()


def parse_table_block(lines, start):
    """从 lines[start] 开始解析表格，返回 (rows, next_index)。"""
    rows = []
    i = start
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        line = lines[i].strip()
        # 跳过分隔行 |---|---|
        if re.match(r"^\|[\s:|-]+\|$", line):
            i += 1
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows.append(cells)
        i += 1
    return rows, i


def add_code_block(lines):
    p = doc.add_paragraph()
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    p.paragraph_format.left_indent = Pt(12)


# ---------------------------------------------------------------------------
# 主解析
# ---------------------------------------------------------------------------
lines = SRC.read_text(encoding="utf-8").splitlines()
i = 0
in_code = False
code_buf = []

while i < len(lines):
    line = lines[i]

    # 代码块
    if line.strip().startswith("```"):
        if in_code:
            add_code_block(code_buf)
            code_buf = []
            in_code = False
        else:
            in_code = True
        i += 1
        continue
    if in_code:
        code_buf.append(line)
        i += 1
        continue

    stripped = line.strip()

    # 表格
    if stripped.startswith("|"):
        rows, i = parse_table_block(lines, i)
        flush_table(rows)
        continue

    # 标题
    if stripped.startswith("### "):
        doc.add_heading(stripped[4:], level=3)
        i += 1
        continue
    if stripped.startswith("## "):
        doc.add_heading(stripped[3:], level=2)
        i += 1
        continue
    if stripped.startswith("# "):
        doc.add_heading(stripped[2:], level=1)
        i += 1
        continue

    # 分割线
    if stripped in ("---", "***"):
        i += 1
        continue

    # 空行
    if not stripped:
        i += 1
        continue

    # 引用
    if stripped.startswith("> "):
        p = add_rich_paragraph(stripped[2:])
        p.paragraph_format.left_indent = Pt(18)
        i += 1
        continue

    # 无序列表
    if stripped.startswith("- "):
        add_rich_paragraph(stripped[2:], bullet=True)
        i += 1
        continue

    # 普通段落
    add_rich_paragraph(stripped)
    i += 1

doc.save(str(DST))
print(f"OK: {DST}")
