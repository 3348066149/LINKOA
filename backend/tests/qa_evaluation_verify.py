# -*- coding: utf-8 -*-
r"""客户满意度评价（V1.4.17）独立验收脚本（QA，Edward）。

绕过 HTTP 认证，直接用临时 SQLite 文件库调用 service 层函数，构造 dict 模拟 current_user，
实测三条红线 + 核心业务行为：

  1. 提交成功落库（三项分数、备注、evaluate_user_id、evaluate_time）；
  2. 同单位客户 A 提交后，客户 B 再提交 → 一次性锁定（已评价，不可重复提交）；
  3. 非「已归档」项目提交 → 报「尚未归档」；
  4. 客户提交非本单位项目（category_id 不匹配）→ 报「无权评价」；
  5. 客户 get 返回不含 reflection；内部 get 返回含 reflection（反思客户不可见红线）；
  6. 无评价时写反思 → 「客户尚未评价」；有评价后写反思覆盖成功；
  7. 消息通知：create_user + 设计师 + 全部管理员，msg_type=evaluation_create；
  8. Excel 导出追加 6 列，已评价有值 / 未评价留空；
  9. （回归）project_list 应能正常返回——顺带暴露 _evaluation_map 是否缺失。

用法：
  cd backend && ..\.venv\Scripts\python.exe tests\qa_evaluation_verify.py

不触碰真实 linkoa.db；脚本结束自动删除临时库文件。
"""
import os
import sys
import tempfile

# ---- 1. 先用临时库路径覆盖 DB_URL，再导入 app（避免触碰真实 linkoa.db）----
TMP_DB = os.path.join(tempfile.gettempdir(), f"qa_eval_verify_{os.getpid()}.db")
if os.path.exists(TMP_DB):
    os.remove(TMP_DB)
os.environ["DB_URL"] = f"sqlite:///{TMP_DB}"

BACKEND_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

from sqlalchemy import select  # noqa: E402

from app.database import Base, SessionLocal, engine, prepare_sqlite_indexes  # noqa: E402
import app.models  # noqa: E402,F401  触发全量模型注册
from app.models import (  # noqa: E402
    CustomerCategory,
    Customer,
    Message,
    Project,
    ProjectDesignerRel,
    ProjectEvaluation,
    SysUser,
)
from app.schemas import EvaluationSubmitRequest, ReflectionSaveRequest  # noqa: E402
from app.services import evaluation_service, excel_service, project_service  # noqa: E402
from app.utils.exceptions import BusinessException  # noqa: E402

PASSED = []
FAILED = []


def check(name: str, cond, detail: str = ""):
    if cond:
        PASSED.append(name)
        print(f"  [PASS] {name}")
    else:
        FAILED.append((name, detail or "断言不成立"))
        print(f"  [FAIL] {name}: {detail or '断言不成立'}")


def expect_business(name: str, fn, code: int, msg_contains: str):
    """断言 fn() 抛出 BusinessException(code, msg 含指定子串)"""
    try:
        fn()
        check(name, False, f"未抛出 BusinessException（期望 {code}: {msg_contains}）")
    except BusinessException as e:
        ok_code = e.code == code
        ok_msg = msg_contains in e.msg
        check(name, ok_code and ok_msg, f"code={e.code}(期望{code}) msg={e.msg!r}(期望含{msg_contains!r})")
    except Exception as e:  # noqa: BLE001
        check(name, False, f"抛出非 BusinessException: {type(e).__name__}: {e}")


def main() -> int:
    # ---- 2. 建全表（临时库）----
    # SQLite 索引名须全库唯一，与 init_database 一致：先给索引名加表名前缀
    prepare_sqlite_indexes()
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()

    try:
        # ---- 3. 造数据：单位、客户、用户、项目 ----
        cat_a = CustomerCategory(category_name="单位A")
        cat_b = CustomerCategory(category_name="单位B")
        db.add_all([cat_a, cat_b])
        db.flush()

        admin = SysUser(username="admin_qa", password="x", real_name="管理员QA", role="管理员", is_enable=1)
        designer = SysUser(username="designer_qa", password="x", real_name="设计师QA", role="普通", is_enable=1)
        cust_a = SysUser(username="cust_a", password="x", real_name="客户A", role="客户", category_id=cat_a.id, is_enable=1)
        cust_b = SysUser(username="cust_b", password="x", real_name="客户B", role="客户", category_id=cat_a.id, is_enable=1)
        cust_other = SysUser(username="cust_o", password="x", real_name="客户O", role="客户", category_id=cat_b.id, is_enable=1)
        db.add_all([admin, designer, cust_a, cust_b, cust_other])
        db.flush()

        cust_ent_a = Customer(category_id=cat_a.id, customer_name="对接人A")
        cust_ent_b = Customer(category_id=cat_b.id, customer_name="对接人B")
        db.add_all([cust_ent_a, cust_ent_b])
        db.flush()

        # 已归档、属于单位A 的项目（用于主流程）
        p_archived = Project(
            project_name="已归档项目A", category_id=cat_a.id, customer_id=cust_ent_a.id,
            project_status="已归档", create_user_id=admin.id,
        )
        # 进行中、属于单位A 的项目（用于「尚未归档」校验）
        p_ongoing = Project(
            project_name="进行中项目A", category_id=cat_a.id, customer_id=cust_ent_a.id,
            project_status="进行中", create_user_id=admin.id,
        )
        # 已归档、属于单位B 的项目（用于「无权评价」隔离校验）
        p_other = Project(
            project_name="单位B已归档", category_id=cat_b.id, customer_id=cust_ent_b.id,
            project_status="已归档", create_user_id=admin.id,
        )
        # 无评价的归档项目（用于「客户尚未评价」写反思校验）
        p_noeval = Project(
            project_name="无评价归档项目", category_id=cat_a.id, customer_id=cust_ent_a.id,
            project_status="已归档", create_user_id=admin.id,
        )
        db.add_all([p_archived, p_ongoing, p_other, p_noeval])
        db.flush()
        db.add(ProjectDesignerRel(project_id=p_archived.id, user_id=designer.id))
        db.commit()

        uid_a = cust_a.id
        uid_b = cust_b.id
        admin_id = admin.id

        # ---- 4. 模型层静态检查 ----
        meta = Base.metadata.tables["project_evaluation"]
        has_unique = any(getattr(idx, "unique", False) for idx in meta.indexes)
        check("model.project_id 唯一索引", has_unique, f"索引: {[i.name for i in meta.indexes]}")

        # ---- 5. 主流程：客户 A 提交成功落库 ----
        req = EvaluationSubmitRequest(
            project_id=p_archived.id, design_score=5, service_score=4, product_score=3,
            customer_remark="整体不错",
        )
        out = evaluation_service.submit(
            db, {"id": uid_a, "role": "客户", "category_id": cat_a.id}, req
        )
        check("submit 返回公开出参", out.design_score == 5 and out.service_score == 4 and out.product_score == 3)
        check("submit 返回不含 reflection", not hasattr(out, "reflection"))

        ev = db.execute(
            select(ProjectEvaluation).where(ProjectEvaluation.project_id == p_archived.id)
        ).scalars().first()
        check("落库三项分数", ev and (ev.design_score, ev.service_score, ev.product_score) == (5, 4, 3))
        check("落库备注", ev is not None and ev.customer_remark == "整体不错")
        check("落库 evaluate_user_id=客户A", ev is not None and ev.evaluate_user_id == uid_a)
        check("落库 evaluate_time", ev is not None and ev.evaluate_time is not None)

        # ---- 6. 一次性锁定：客户 B 再提交 ----
        req2 = EvaluationSubmitRequest(
            project_id=p_archived.id, design_score=1, service_score=1, product_score=1,
        )
        expect_business(
            "同单位客户B重复提交被拒",
            lambda: evaluation_service.submit(db, {"id": uid_b, "role": "客户", "category_id": cat_a.id}, req2),
            400, "已评价",
        )

        # ---- 7. 非归档项目提交 ----
        expect_business(
            "非归档项目提交被拒",
            lambda: evaluation_service.submit(
                db, {"id": uid_a, "role": "客户", "category_id": cat_a.id},
                EvaluationSubmitRequest(project_id=p_ongoing.id, design_score=5, service_score=5, product_score=5),
            ),
            400, "尚未归档",
        )

        # ---- 8. 非本单位项目提交（隔离红线）----
        expect_business(
            "客户提交非本单位项目被拒",
            lambda: evaluation_service.submit(
                db, {"id": uid_a, "role": "客户", "category_id": cat_a.id},
                EvaluationSubmitRequest(project_id=p_other.id, design_score=5, service_score=5, product_score=5),
            ),
            403, "无权评价",
        )

        # ---- 9. 客户 get 无 reflection / 内部 get 含 reflection ----
        pub = evaluation_service.get(db, p_archived.id, {"id": uid_a, "role": "客户", "category_id": cat_a.id})
        check("客户 get 返回公开 DTO（无 reflection）", pub is not None and not hasattr(pub, "reflection"))
        detail = evaluation_service.get(db, p_archived.id, {"id": admin_id, "role": "管理员"})
        check("内部 get 返回 Detail DTO（含 reflection）", detail is not None and hasattr(detail, "reflection"))

        # 客户查看非本单位项目 → 403
        expect_business(
            "客户 get 非本单位项目被拒",
            lambda: evaluation_service.get(db, p_other.id, {"id": uid_a, "role": "客户", "category_id": cat_a.id}),
            403, "无权查看",
        )

        # ---- 10. 反思：无评价报错，有评价覆盖成功 ----
        expect_business(
            "无评价写反思被拒",
            lambda: evaluation_service.save_reflection(
                db, {"id": admin_id, "role": "管理员"},
                ReflectionSaveRequest(project_id=p_noeval.id, reflection="复盘"),
            ),
            400, "客户尚未评价",
        )

        evaluation_service.save_reflection(
            db, {"id": admin_id, "role": "管理员"},
            ReflectionSaveRequest(project_id=p_archived.id, reflection="第一次反思"),
        )
        d1 = evaluation_service.get(db, p_archived.id, {"id": admin_id, "role": "管理员"})
        check("写反思成功", d1 is not None and d1.reflection == "第一次反思")
        check("反思落库编辑人姓名", d1 is not None and d1.reflection_user_name == admin.real_name)
        check("反思落库 reflection_time", d1 is not None and d1.reflection_time is not None)

        evaluation_service.save_reflection(
            db, {"id": designer.id, "role": "普通"},
            ReflectionSaveRequest(project_id=p_archived.id, reflection="覆盖后的反思"),
        )
        d2 = evaluation_service.get(db, p_archived.id, {"id": admin_id, "role": "管理员"})
        check("反思覆盖式更新", d2 is not None and d2.reflection == "覆盖后的反思")
        check("反思覆盖更新编辑人", d2 is not None and d2.reflection_user_name == designer.real_name)

        # ---- 11. 消息通知：创建人 + 设计师 + 全部管理员 ----
        msgs = db.execute(
            select(Message).where(Message.project_id == p_archived.id, Message.msg_type == "evaluation_create")
        ).scalars().all()
        recv = sorted(m.receive_user_id for m in msgs)
        expected = sorted({admin_id, designer.id})
        check("消息接收人=创建人+设计师+管理员(去重)", recv == expected, f"recv={recv} expected={expected}")
        check("消息内容含项目名", all("已归档项目A" in m.msg_content for m in msgs))

        # ---- 12. Excel 导出追加 6 列 ----
        from openpyxl import load_workbook
        from io import BytesIO
        from app.schemas import ReportExportQuery

        check("HEADERS 共 13 列", len(excel_service.HEADERS) == 13, f"实际 {len(excel_service.HEADERS)}")
        check("COL_WIDTHS 同步 13 列", len(excel_service.COL_WIDTHS) == 13, f"实际 {len(excel_service.COL_WIDTHS)}")
        check(
            "HEADERS 含 6 个新增列名",
            excel_service.HEADERS[-6:] == ["设计满意度", "服务响应速度满意度", "最终产品满意度", "客户备注", "评价时间", "反思/总结"],
            f"实际 {excel_service.HEADERS[-6:]}",
        )

        blob, _ = excel_service.export_projects(db, ReportExportQuery())
        wb = load_workbook(BytesIO(blob))
        ws = wb["项目列表"]
        rows_by_name = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}
        r_arch = rows_by_name.get("已归档项目A")
        r_noeval = rows_by_name.get("无评价归档项目")
        if r_arch:
            check("Excel 已评价行三项评分为数字", ws.cell(row=r_arch, column=8).value == 5 and ws.cell(row=r_arch, column=9).value == 4 and ws.cell(row=r_arch, column=10).value == 3)
            check("Excel 已评价行备注有值", ws.cell(row=r_arch, column=11).value == "整体不错")
            check("Excel 已评价行反思有值", ws.cell(row=r_arch, column=13).value == "覆盖后的反思")
        else:
            check("Excel 含已归档项目A行", False, "未找到行")
        if r_noeval:
            check("Excel 未评价行留空", ws.cell(row=r_noeval, column=8).value in (None, "") and ws.cell(row=r_noeval, column=13).value in (None, ""))
        else:
            check("Excel 含无评价项目行", False, "未找到行")

        # ---- 13. 回归：project_list 不应崩溃（暴露 _evaluation_map 缺失）----
        from app.schemas import ProjectListQuery
        try:
            result = project_service.project_list(
                db, ProjectListQuery(page_num=1, page_size=20),
                {"id": admin_id, "role": "管理员"},
            )
            check("project_list 正常返回", True, "")
        except Exception as e:  # noqa: BLE001
            check("project_list 正常返回", False, f"抛出 {type(e).__name__}: {e}")

    finally:
        db.close()

    # ---- 汇总 ----
    print("\n" + "=" * 60)
    print(f"通过 {len(PASSED)} 项，失败 {len(FAILED)} 项")
    for name, err in FAILED:
        print(f"  FAIL {name}: {err}")
    # 清理临时库
    try:
        os.remove(TMP_DB)
    except OSError:
        pass
    return 0 if not FAILED else 2


if __name__ == "__main__":
    sys.exit(main())
