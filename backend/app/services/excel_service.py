"""报表导出 Service：项目列表 Excel（OpenPyXL）"""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import CustomerCategory, Project, ProjectEvaluation
from app.schemas import ReportExportQuery
from app.services.project_service import _contact_names, _designer_names

HEADERS = [
    "项目名称", "客户单位", "客户对接人", "参与设计师", "立项时间", "预计结束时间", "项目状态",
    "设计满意度", "服务响应速度满意度", "最终产品满意度", "客户备注", "评价时间", "反思/总结",
]
COL_WIDTHS = [30, 20, 20, 30, 18, 18, 12, 12, 16, 14, 30, 18, 30]


def _evaluation_export_map(db: Session, project_ids: list[int]) -> dict[int, ProjectEvaluation]:
    """批量取项目评价（含反思），一次 IN 查询防 N+1"""
    if not project_ids:
        return {}
    rows = db.execute(
        select(ProjectEvaluation).where(ProjectEvaluation.project_id.in_(project_ids))
    ).scalars().all()
    return {e.project_id: e for e in rows}


def export_projects(db: Session, query: ReportExportQuery) -> tuple[bytes, str]:
    """导出项目列表 Excel，返回 (文件字节流, 文件名)"""
    stmt = (
        select(Project, CustomerCategory.category_name.label("company_name"))
        .outerjoin(CustomerCategory, CustomerCategory.id == Project.category_id)
        .order_by(Project.id.desc())
    )
    if query.project_status:
        stmt = stmt.where(Project.project_status == query.project_status)
    if query.category_id:
        stmt = stmt.where(Project.category_id == query.category_id)
    if query.start_time:
        stmt = stmt.where(Project.create_time >= query.start_time)
    if query.end_time:
        stmt = stmt.where(Project.create_time <= query.end_time)

    rows = db.execute(stmt).all()
    name_map = _designer_names(db, [p.id for p, _ in rows])
    contact_map = _contact_names(db, [p.id for p, _ in rows])
    eval_map = _evaluation_export_map(db, [p.id for p, _ in rows])

    wb = Workbook()
    ws = wb.active
    ws.title = "项目列表"

    # 表头：加粗 + 背景色
    header_font = Font(bold=True)
    header_fill = PatternFill(fill_type="solid", fgColor="D9E1F2")
    for col, title in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        ws.column_dimensions[cell.column_letter].width = COL_WIDTHS[col - 1]

    # 数据行（已停用用户参与的历史项目正常统计，不过滤）
    for row_idx, (p, company_name) in enumerate(rows, start=2):
        ws.cell(row=row_idx, column=1, value=p.project_name)
        ws.cell(row=row_idx, column=2, value=company_name or "")
        ws.cell(row=row_idx, column=3, value=contact_map.get(p.id, ""))
        ws.cell(row=row_idx, column=4, value=name_map.get(p.id, ""))
        ws.cell(row=row_idx, column=5, value=p.create_time.strftime("%Y-%m-%d %H:%M") if p.create_time else "")
        ws.cell(row=row_idx, column=6, value=p.expect_time.strftime("%Y-%m-%d %H:%M") if p.expect_time else "")
        ws.cell(row=row_idx, column=7, value=p.project_status)
        ev = eval_map.get(p.id)
        ws.cell(row=row_idx, column=8, value=ev.design_score if ev else "")
        ws.cell(row=row_idx, column=9, value=ev.service_score if ev else "")
        ws.cell(row=row_idx, column=10, value=ev.product_score if ev else "")
        ws.cell(row=row_idx, column=11, value=ev.customer_remark if ev else "")
        ws.cell(row=row_idx, column=12, value=ev.evaluate_time.strftime("%Y-%m-%d %H:%M") if ev and ev.evaluate_time else "")
        ws.cell(row=row_idx, column=13, value=ev.reflection if ev else "")

    buf = io.BytesIO()
    wb.save(buf)
    filename = f"项目列表_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return buf.getvalue(), filename
